# -*- coding: utf-8 -*-
"""엑셀 DB(MMM_DB)의 Balance_Log 시트에 밸런스 실험 결과를 한 줄씩 기록한다. 쓰는 사람은 관리자(메인 에이전트) 하나.
사용: python BalanceData/db/db_log.py <결과.json>   (같은 Exp_ID가 있으면 그 줄을 덮어쓴다)
결과.json: {"Exp_ID": "T1", ...아래 COLS 키...}"""
import io, json, os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill

DB = os.environ.get("MMM_DB", r"C:\Users\user\Desktop\MMM_DB_Ver.02.xlsx")
SHEET = "Balance_Log"
# (한글 이름, 키, 타입)
COLS = [
    ("실험 ID", "Exp_ID", "string"), ("일시", "Exp_Time", "string"), ("담당", "Exp_By", "string"),
    ("프리셋", "Preset", "string"), ("가격표", "Price_Ver", "string"), ("시드", "Seeds", "string"), ("바 오차", "Skill", "float"),
    ("트리 완료 판", "Done_Round", "string"), ("트리 완료 분", "Done_Min", "string"),
    ("초반 완료", "Done_T1", "string"), ("중반 완료", "Done_T2", "string"), ("후반 완료", "Done_T3", "string"),
    ("산 판 비율", "Buy_Rate", "string"), ("구매 간격 중앙/최대", "Buy_Gap", "string"),
    ("가격 최대 뜀", "Step_Max", "float"), ("노드 안 최대", "Node_Step", "float"), ("효율 넘김", "Cross", "string"),
    ("판정", "Verdict", "string"), ("관리자 검증", "Checked", "string"), ("출처", "Source", "string"), ("비고", "Note", "string"),
]


def sheet(wb):
    if SHEET in wb.sheetnames:
        return wb[SHEET]
    ws = wb.create_sheet(SHEET)
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="DDEBF7")
    for c, (kr, key, typ) in enumerate(COLS, 1):
        for r, v in enumerate((kr, key, typ), 1):
            cell = ws.cell(r, c, v)
            cell.font = bold
            if r < 3:
                cell.fill = fill
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = max(10, min(40, len(kr) * 2 + 4))
    ws.freeze_panes = "B4"
    return ws


def write(rec):
    wb = openpyxl.load_workbook(DB)
    ws = sheet(wb)
    keys = [k for _, k, _ in COLS]
    row = next((r for r in range(4, ws.max_row + 1) if ws.cell(r, 1).value == rec["Exp_ID"]), None) or ws.max_row + 1
    if row < 4:
        row = 4
    for c, k in enumerate(keys, 1):
        v = rec.get(k, 0)
        ws.cell(row, c, v if v is not None else 0)
    wb.save(DB)
    return row


if __name__ == "__main__":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    rec = json.load(open(sys.argv[1], encoding="utf-8"))
    print(f"{SHEET} {write(rec)}행: {rec['Exp_ID']}")
