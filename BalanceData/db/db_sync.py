# -*- coding: utf-8 -*-
"""엑셀 DB(MMM_DB)의 Upgrade · Upgrade_Cost 시트를 지금 게임 코드(nodes_export.json) 기준으로 맞춘다.
다시 돌릴 수 있다 (같은 입력 → 같은 결과). 입력 파일은 읽기만 하고, 결과는 다른 파일로 낸다.

사용: python BalanceData/db/db_sync.py <입력 xlsx> <출력 xlsx>
  예) python BalanceData/db/db_sync.py C:\\Users\\user\\Desktop\\MMM_DB_Ver.02.xlsx BalanceData/db/MMM_DB_staged.xlsx

하는 일
  1. 코드 노드 <-> DB 행 대응: code_map.json({code_id: Upg_ID})이 있으면 먼저 쓰고, 나머지는 이름(Upg_Name == name)이
     같고 ID 접두(11 ps / 12 md / 13 ed)가 맞는 행끼리만 잇는다. 이름이 바뀐 듯한 행은 추정하지 않고 '미대응'.
  2. 대응된 행: Upg_Name · Upg_Pre · Efx_MaxLv · Efx_Lv1~4 · Cost_ID 를 코드 값으로 (Type · Str_Group_ID 는 유지).
  3. 코드에만 있는 노드: 새 행 (Upg_ID = 접두 다음 빈 번호, Str_Group_ID = 2nnnn 다음 번호, Type = 같은 계열 이름의 기존 행을 따름, 모르면 0).
  4. 4레벨: Upgrade K열 = Efx_Lv4, Upgrade_Cost E열 = Cost_Lv4 (없는 레벨은 0).
  5. 값: Efx = 레벨 value. value가 null이면 text에서 규칙대로: 코어(core_*) = 1 + %/100 · 해금형(max 1이고 skill 노드이거나
     문구가 '사용 중'/'열림'으로 시작) = 1 · *_rate = %/100 · *_pow = 레벨 번호 · 그 밖에 %가 있으면 %/100, 없으면 첫 숫자.
     Cost = 레벨 gold. 노드마다 자기 Cost_ID.
  (엑셀에 쓰는 일은 compute()가 계산한 결과를 쓴다. 이 스크립트는 openpyxl로 저장하므로 Skima 도형이 사라진다 ->
   실제 적용은 db_apply_com.py (Excel COM).)
  6. DB에만 있는 행은 그대로 둔다 (목록만 출력).
  7. Info 시트의 '(1~3)' 문구 3곳을 4레벨로. 다른 시트는 그대로.
"""
import io, json, os, re, sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
NODES = os.path.join(HERE, "nodes_export.json")
CODE_MAP = os.path.join(HERE, "code_map.json")

PREFIX = {"ps": 11, "md": 12, "ed": 13}            # br -> Upgrade ID 접두 (Info 규칙)
EXTRA_PREFIX = 14                                   # br가 위 셋 밖일 때 (gm 등)
STR_BASE = 10000                                    # 업그레이드 Str_Group_ID = Upg_ID + 10000 대역 (21/22/23nnn)
COST_BASE = 90000
MAXLV = 4
DATA_ROW = 4                                        # 4행부터 데이터


def out(*a):
    print(*a)


def num(v):
    """float을 깔끔하게 (정수면 int, 아니면 소수 6자리)."""
    v = round(float(v), 6)
    return int(v) if v == int(v) else v


def first_number(text):
    m = re.search(r"\d+(?:\.\d+)?", text or "")
    return float(m.group()) if m else None


def family(name):
    """이름 계열: 괄호 · 끝의 '+n' · 로마 숫자를 뗀 이름. (타격 점수 II+ -> 타격 점수, 핀볼 +1 (II) -> 핀볼)"""
    n = re.sub(r"\s*\([^)]*\)", "", name)
    n = re.sub(r"\s*\+\d*$", "", n)
    n = re.sub(r"\s+(?:I{1,3}|IV|V)$", "", n)
    return n.strip()


def pct(text):
    m = re.search(r"(\d+(?:\.\d+)?)\s*%", text or "")
    return float(m.group(1)) if m else None


def text_value(node, lv):
    """value가 null인 레벨의 값. 규칙은 모듈 설명 5번."""
    text = lv["text"] or ""
    nid = node["id"]
    if node.get("core") or nid.startswith("core_"):
        p = pct(text)
        return 1 + p / 100 if p is not None else 1
    if nid.endswith("_rate"):
        p = pct(text)
        return p / 100 if p is not None else first_number(text)
    if nid.endswith("_pow"):
        return lv["L"]
    if node["max"] == 1 and (node.get("skill") or text.startswith("사용 중") or text.startswith("열림")):
        return 1
    p = pct(text)
    if p is not None:
        return p / 100
    f = first_number(text)
    return f if f is not None else 1


def level_values(node):
    """[(Efx, Cost, from_text)] x MAXLV, 없는 레벨은 (0, 0, False)."""
    res = []
    for L in range(1, MAXLV + 1):
        lv = next((x for x in node["levels"] if x["L"] == L), None)
        if lv is None:
            res.append((0, 0, False))
            continue
        fromtext = lv["value"] is None
        v = text_value(node, lv) if fromtext else lv["value"]
        res.append((num(v), int(round(lv["gold"] or 0)), fromtext))
    return res


def compute(src):
    """src(읽기만)를 코드와 맞춘 결과를 계산한다. 파일은 쓰지 않는다."""
    nodes = json.load(io.open(NODES, encoding="utf-8"))["nodes"]
    code_map = json.load(io.open(CODE_MAP, encoding="utf-8")) if os.path.exists(CODE_MAP) else {}

    wb = openpyxl.load_workbook(src)
    ws = wb["Upgrade"]
    wc = wb["Upgrade_Cost"]
    ref_meaning = {}                 # Type -> 값 의미 (Upgrade L열 참고표 M/P)
    for r in range(DATA_ROW, ws.max_row + 1):
        if isinstance(ws.cell(r, 13).value, str):
            ref_meaning[ws.cell(r, 13).value] = ws.cell(r, 16).value or ""

    # ---- DB 읽기 ----
    old_max_row = ws.max_row
    rows = []        # 각 행: dict(id,name,str,pre,type,max,cost,lv[1..3], lv4)
    for r in range(DATA_ROW, ws.max_row + 1):
        if not isinstance(ws.cell(r, 1).value, int):
            continue
        v = [ws.cell(r, c).value for c in range(1, 12)]
        rows.append(dict(id=v[0], name=v[1], str=v[2], pre=v[3], type=v[4], max=v[5], cost=v[6],
                         lv=[v[7], v[8], v[9], 0 if v[10] is None else v[10]], src="db"))
    old_cost_max_row = wc.max_row
    cost_rows = {}
    for r in range(DATA_ROW, wc.max_row + 1):
        cid = wc.cell(r, 1).value
        if isinstance(cid, int):
            cost_rows[cid] = [wc.cell(r, c).value or 0 for c in range(2, 6)]
    by_id = {x["id"]: x for x in rows}

    # ---- 1. 대응 ----
    assign = {}                      # code_id -> Upg_ID
    claimed = set()
    notes = []
    for n in nodes:
        uid = code_map.get(n["id"])
        if uid is not None and uid not in claimed:
            assign[n["id"]] = uid
            claimed.add(uid)
    name_rows = {}
    for x in rows:
        name_rows.setdefault((x["id"] // 1000, x["name"]), []).append(x)
    for n in nodes:
        if n["id"] in assign:
            continue
        pre = PREFIX.get(n["br"], EXTRA_PREFIX)
        cand = [x for x in name_rows.get((pre, n["name"]), []) if x["id"] not in claimed]
        if len(cand) == 1:
            assign[n["id"]] = cand[0]["id"]
            claimed.add(cand[0]["id"])
        elif len(cand) > 1:
            notes.append("이름 중복으로 미대응: %s (%s) <- %s" % (n["id"], n["name"], [c["id"] for c in cand]))
        else:
            other = [x["id"] for k, v in name_rows.items() if k[1] == n["name"] for x in v]
            if other:
                notes.append("이름은 같지만 접두가 달라 미대응: %s (%s, br=%s) <- DB %s" % (n["id"], n["name"], n["br"], other))
    matched_ids = set(claimed & set(by_id))
    # 새 행 ID
    used_max = {}
    for x in rows:
        used_max[x["id"] // 1000] = max(used_max.get(x["id"] // 1000, 0), x["id"])
    for uid in assign.values():
        used_max[uid // 1000] = max(used_max.get(uid // 1000, 0), uid)
    for n in [n for n in nodes if n["id"] not in assign]:
        pre = PREFIX.get(n["br"], EXTRA_PREFIX)
        nid = max(used_max.get(pre, pre * 1000), pre * 1000) + 1
        used_max[pre] = nid
        assign[n["id"]] = nid
    new_nodes = [n for n in nodes if assign[n["id"]] not in by_id]          # DB에 행이 없는 노드 (code_map이 ID만 정해 둔 것 포함)
    new_ids = {assign[n["id"]] for n in new_nodes}
    node_by_uid = {assign[n["id"]]: n for n in nodes}
    code_by_id = {n["id"]: n for n in nodes}

    # ---- Type / Str_Group_ID 근거 (새 행) ----
    fam_type = {}
    for x in sorted(rows, key=lambda x: x["id"]):
        if x["type"] not in (0, None, "0"):
            fam_type.setdefault((x["id"] // 1000, family(x["name"])), x)
    for n in nodes:                  # 대응된 코드 노드의 이름도 계열로 인정 (이름이 바뀐 행의 Type을 계열이 따른다)
        x = by_id.get(assign[n["id"]])
        if x is not None and x["type"] not in (0, None, "0"):
            fam_type.setdefault((x["id"] // 1000, family(n["name"])), x)
    str_max = {}
    for x in rows:
        if isinstance(x["str"], int):
            str_max[x["str"] // 1000] = max(str_max.get(x["str"] // 1000, 0), x["str"])
    type_log = []
    new_rows = []
    for n in new_nodes:
        uid = assign[n["id"]]
        pre = uid // 1000
        sp = pre + 10
        str_max[sp] = max(str_max.get(sp, sp * 1000), sp * 1000) + 1
        ref = fam_type.get((pre, family(n["name"])))
        if ref is not None:
            typ, why = ref["type"], "같은 계열 '%s'(%d)의 Type" % (ref["name"], ref["id"])
            if family(ref["name"]) != family(n["name"]):
                why += " (코드 이름 기준 계열)"
        else:
            typ, why = 0, "같은 계열 이름의 기존 행 없음 -> 0"
        type_log.append((n["id"], uid, typ, why))
        new_rows.append(dict(id=uid, name=n["name"], str=str_max[sp], pre=0, type=typ, max=0, cost=0,
                             lv=[0, 0, 0, 0], src="new"))
    for x in new_rows:
        by_id[x["id"]] = x
    rows += new_rows

    # ---- 2/5. 값 갱신 + Cost_ID ----
    # 원래 Cost_ID 사용처: DB에만 있는 행이 쓰는 ID는 그 행을 위해 값을 남긴다.
    orphan_ids = {x["id"] for x in rows if x["src"] == "db" and x["id"] not in node_by_uid}
    cost_users = {}
    for x in rows:
        if x["src"] == "db":
            cost_users.setdefault(x["cost"], []).append(x["id"])
    frozen_cost = {x["cost"] for x in rows if x["id"] in orphan_ids}
    next_cost = max([COST_BASE] + list(cost_rows.keys())) + 1
    taken_cost = set()
    node_cost = {}                   # Upg_ID -> Cost_ID
    cost_changed = []
    for x in sorted(rows, key=lambda x: (x["src"] == "new", x["id"])):
        if x["id"] not in node_by_uid:
            continue
        old = x["cost"] if x["src"] == "db" else None
        if old is not None and old not in frozen_cost and old not in taken_cost:
            cid = old
        else:
            cid = next_cost
            next_cost += 1
            cost_changed.append((node_by_uid[x["id"]]["id"], x["id"], old, cid))
        taken_cost.add(cid)
        node_cost[x["id"]] = cid

    text_log = []
    final_cost = {}                  # Cost_ID -> [lv1..lv4]
    for x in rows:
        n = node_by_uid.get(x["id"])
        if n is None:
            continue
        vals = level_values(n)
        x["name"] = n["name"]
        x["pre"] = assign[n["parent"]] if n["parent"] else 0
        x["max"] = n["max"]
        x["cost"] = node_cost[x["id"]]
        x["lv"] = [v[0] for v in vals]
        final_cost[x["cost"]] = [v[1] for v in vals]
        for L, v in enumerate(vals, 1):
            if v[2]:
                text_log.append((n["id"], L, v[0], next(l["text"] for l in n["levels"] if l["L"] == L)))
    for x in rows:                   # DB에만 있는 행이 쓰는 비용: 원래 값 유지
        if x["id"] not in node_by_uid and x["cost"] not in final_cost:
            final_cost[x["cost"]] = list(cost_rows.get(x["cost"], [0, 0, 0, 0]))[:4]
    rows.sort(key=lambda x: x["id"])

    # 값 의미(배율 x)인 Type인데 코드 문구는 더하기(+)인 노드 -> 보고용
    meaning_changes = []
    for x in rows:
        n = node_by_uid.get(x["id"])
        if n is None or x["type"] in (0, None, "0"):
            continue
        if "배율" in ref_meaning.get(x["type"], "") and all(("+" in l["text"] and "×" not in l["text"] and "%" not in l["text"]) for l in n["levels"]):
            meaning_changes.append((n["id"], x["id"], x["type"], n["levels"][0]["text"]))

    used_costs = sorted({x["cost"] for x in rows})
    info_changes = []                # (시트 셀 주소, 옛 문구, 새 문구)
    for row in wb["Info"].iter_rows():
        for cell in row:
            v = cell.value
            if not isinstance(v, str):
                continue
            nv = v
            if v.strip() == "Efx_Lv1~3":
                nv = "Efx_Lv1~4"
            elif v.strip() == "Cost_Lv1~3":
                nv = "Cost_Lv1~4"
            elif "강화 최대 레벨" in v and "(1~3)" in v:
                nv = v.replace("(1~3)", "(1~4)")
            if nv != v:
                info_changes.append((cell.coordinate, v, nv))
    dbonly = [x for x in rows if x["id"] not in node_by_uid]
    return dict(nodes=nodes, assign=assign, rows=rows, final_cost=final_cost, used_costs=used_costs,
                new_nodes=new_nodes, new_ids=new_ids, new_rows=new_rows, type_log=type_log, dbonly=dbonly,
                text_log=text_log, cost_changed=cost_changed, info_changes=info_changes, notes=notes,
                meaning_changes=meaning_changes, node_by_uid=node_by_uid,
                old_max_row=old_max_row, old_cost_max_row=old_cost_max_row)


def main(src, dst):
    if os.path.abspath(src) == os.path.abspath(dst):
        sys.exit("입력과 출력이 같은 파일이다. 원본에는 쓰지 않는다.")
    R = compute(src)
    nodes, assign, rows, final_cost = R["nodes"], R["assign"], R["rows"], R["final_cost"]
    used_costs, new_nodes, new_ids = R["used_costs"], R["new_nodes"], R["new_ids"]
    new_rows, type_log, dbonly, text_log = R["new_rows"], R["type_log"], R["dbonly"], R["text_log"]
    cost_changed, info_changes, notes = R["cost_changed"], R["info_changes"], R["notes"]
    old_max_row, old_cost_max_row = R["old_max_row"], R["old_cost_max_row"]
    wb = openpyxl.load_workbook(src)
    ws = wb["Upgrade"]
    wc = wb["Upgrade_Cost"]

    # ---- 시트 쓰기: Upgrade (A~K) ----
    tmpl = {c: ws.cell(DATA_ROW, c)._style for c in range(1, 11)}
    tmpl[11] = tmpl[10]
    for c, (a, b, d) in ((11, ("4레벨", "Efx_Lv4", "float")),):
        for r, val in ((1, a), (2, b), (3, d)):
            ws.cell(r, c).value = val
            ws.cell(r, c)._style = ws.cell(r, 10)._style
    ws.column_dimensions["K"].width = ws.column_dimensions["J"].width or 9
    for r in range(DATA_ROW, max(old_max_row, DATA_ROW + len(rows)) + 1):
        for c in range(1, 12):
            ws.cell(r, c).value = None
    for i, x in enumerate(rows):
        r = DATA_ROW + i
        vals = [x["id"], x["name"], x["str"], x["pre"], x["type"], x["max"], x["cost"]] + list(x["lv"])
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c)
            cell.value = v
            cell._style = tmpl[c]

    # ---- Upgrade_Cost (A~E) ----
    ctmpl = {c: wc.cell(DATA_ROW, c)._style for c in range(1, 5)}
    ctmpl[5] = ctmpl[4]
    for r, val in ((1, "4레벨"), (2, "Cost_Lv4"), (3, "int")):
        wc.cell(r, 5).value = val
        wc.cell(r, 5)._style = wc.cell(r, 4)._style
    wc.column_dimensions["E"].width = 12
    for r in range(DATA_ROW, max(old_cost_max_row, DATA_ROW + len(used_costs)) + 1):
        for c in range(1, 6):
            wc.cell(r, c).value = None
    for i, cid in enumerate(used_costs):
        r = DATA_ROW + i
        for c, v in enumerate([cid] + [int(t) for t in final_cost[cid]], 1):
            cell = wc.cell(r, c)
            cell.value = v
            cell._style = ctmpl[c]

    for coord, old, new in info_changes:
        wb["Info"][coord].value = new

    wb.save(dst)
    with io.open(CODE_MAP, "w", encoding="utf-8") as f:
        json.dump({n["id"]: assign[n["id"]] for n in nodes}, f, ensure_ascii=False, indent=1)

    # ---- 보고 ----
    out("입력:", src)
    out("출력:", dst)
    out("코드 노드 %d · 대응 %d · 새 행 %d (%s) · DB에만 있는 행 %d · 최종 Upgrade 행 %d · Cost 행 %d"
        % (len(nodes), len(nodes) - len(new_nodes), len(new_nodes),
           ("%d~%d" % (min(new_ids), max(new_ids))) if new_ids else "-", len(dbonly), len(rows), len(used_costs)))
    for s in notes:
        out("[주의]", s)
    out("\n[새 행] code_id | 새 Upg_ID | Str_Group_ID | Type | 근거")
    sg = {x["id"]: x["str"] for x in new_rows}
    for cid, uid, typ, why in type_log:
        out("  %s | %d | %d | %s | %s" % (cid, uid, sg[uid], typ, why))
    out("\n[DB에만 있는 행] Upg_ID | 이름 | 선행 | Cost_ID | 같은 부모의 새 행(이름 바뀜 후보)")
    for x in dbonly:
        sib = [n["id"] + "/" + n["name"] for n in new_nodes
               if n["parent"] and assign[n["parent"]] == x["pre"] and x["pre"] != 0]
        out("  %d | %s | %s | %s | %s" % (x["id"], x["name"], x["pre"], x["cost"], ", ".join(sib)))
    out("\n[값을 text에서 뽑은 레벨] %d건" % len(text_log))
    for cid, L, v, t in text_log:
        out("  %s L%d = %s <- %s" % (cid, L, v, t))
    out("\n[Cost_ID 새로 준 노드] %d건 (code_id, Upg_ID, 원래, 새)" % len(cost_changed))
    for t in cost_changed:
        out("  ", t)
    out("\n[Info 문구 변경]", info_changes)
    out("\n[값 의미가 배율(x) -> 더하기(+)로 바뀐 노드] %d건 (code_id, Upg_ID, Type, 1레벨 문구)" % len(R["meaning_changes"]))
    for t in R["meaning_changes"]:
        out("  ", t)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
